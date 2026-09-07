import base64
import importlib.util
import io
import json
from pathlib import Path, PurePosixPath
import posixpath
import struct
import subprocess
import sys
import tempfile
import unittest
from xml.etree import ElementTree as ET
from zipfile import ZipFile
import zlib


ROOT = Path(__file__).resolve().parents[3]
SCRIPT = ROOT / "scripts" / "html2pptx" / "build_pptx.py"
NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "ct": "http://schemas.openxmlformats.org/package/2006/content-types",
}
JPEG = base64.b64decode(
    "/9j/4AAQSkZJRgABAQAAAQABAAD/2wBDAAgGBgcGBQgHBwcJCQgKDBQNDAsLDBkSEw8UHRofHh0aHBwgJC4nICIsIxwcKDcpLDAxNDQ0Hyc5PTgyPC4zNDL/2wBDAQkJCQwLDBgNDRgyIRwhMjIyMjIyMjIyMjIyMjIyMjIyMjIyMjIyMjIyMjIyMjIyMjIyMjIyMjIyMjIyMjIyMjL/wAARCAAJABADASIAAhEBAxEB/8QAHwAAAQUBAQEBAQEAAAAAAAAAAAECAwQFBgcICQoL/8QAtRAAAgEDAwIEAwUFBAQAAAF9AQIDAAQRBRIhMUEGE1FhByJxFDKBkaEII0KxwRVS0fAkM2JyggkKFhcYGRolJicoKSo0NTY3ODk6Q0RFRkdISUpTVFVWV1hZWmNkZWZnaGlqc3R1dnd4eXqDhIWGh4iJipKTlJWWl5iZmqKjpKWmp6ipqrKztLW2t7i5usLDxMXGx8jJytLT1NXW19jZ2uHi4+Tl5ufo6erx8vP09fb3+Pn6/8QAHwEAAwEBAQEBAQEBAQAAAAAAAAECAwQFBgcICQoL/8QAtREAAgECBAQDBAcFBAQAAQJ3AAECAxEEBSExBhJBUQdhcRMiMoEIFEKRobHBCSMzUvAVYnLRChYkNOEl8RcYGRomJygpKjU2Nzg5OkNERUZHSElKU1RVVldYWVpjZGVmZ2hpanN0dXZ3eHl6goOEhYaHiImKkpOUlZaXmJmaoqOkpaanqKmqsrO0tba3uLm6wsPExcbHyMnK0tPU1dbX2Nna4uPk5ebn6Onq8vP09fb3+Pn6/9oADAMBAAIRAxEAPwDk6KKKk+TP/9k="
)


def png():
    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 16, 9, 8, 2, 0, 0, 0)) + chunk(
        b"IDAT", zlib.compress((b"\0" + bytes((200, 40, 30)) * 16) * 9)
    ) + chunk(b"IEND", b"")


class BuildPptxTest(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.directory = Path(self.temporary.name)
        self.output = self.directory / "输出 & 演示.pptx"
        self.images = {"后页.png": png(), "前页.jpg": JPEG}
        for name, data in self.images.items():
            (self.directory / name).write_bytes(data)
        self.order = ["前页.jpg", "后页.png", "前页.jpg"]
        self.manifest(self.order)

    def manifest(self, slides):
        (self.directory / "manifest.json").write_text(json.dumps({"slides": slides}), encoding="utf-8")

    def build(self, embed=None):
        if embed is None:
            command = [sys.executable, "-S", str(SCRIPT)]
        else:
            wrapper = """
import importlib.abc, runpy, sys
class WithoutOfficeLibraries(importlib.abc.MetaPathFinder):
    def find_spec(self, fullname, path=None, target=None):
        if fullname.split('.')[0] in {'pptx', 'lxml', 'xlsxwriter'}:
            raise ImportError('此导出进程禁止使用 Office 文档库：' + fullname)
sys.meta_path.insert(0, WithoutOfficeLibraries())
script = sys.argv.pop(1)
sys.path.insert(0, str(__import__('pathlib').Path(script).parent))
sys.argv[0] = script
runpy.run_path(script, run_name='__main__')
"""
            command = [sys.executable, "-c", wrapper, str(SCRIPT)]
        command += [str(self.directory), str(self.output)]
        if embed is not None:
            command.append(str(embed))
        return subprocess.run(command, capture_output=True, text=True)

    def assert_archive(self, archive):
        names = archive.namelist()
        self.assertEqual(len(names), len(set(names)))
        types = ET.fromstring(archive.read("[Content_Types].xml"))
        defaults = {node.get("Extension"): node.get("ContentType") for node in types.findall("ct:Default", NS)}
        overrides = {node.get("PartName").lstrip("/"): node.get("ContentType") for node in types.findall("ct:Override", NS)}
        for name in names:
            if name == "[Content_Types].xml":
                continue
            self.assertTrue(name in overrides or name.rsplit(".", 1)[-1] in defaults, name)
            if name.endswith(".xml"):
                part = ET.fromstring(archive.read(name))
                referenced_ids = {
                    value for node in part.iter() for key, value in node.attrib.items()
                    if key.startswith("{" + NS["r"] + "}")
                }
                if referenced_ids:
                    relationship_part = str(PurePosixPath(name).parent / "_rels" / (PurePosixPath(name).name + ".rels"))
                    self.assertIn(relationship_part, names, name)
                    declared_ids = {node.get("Id") for node in ET.fromstring(archive.read(relationship_part))}
                    self.assertTrue(referenced_ids <= declared_ids, (name, referenced_ids - declared_ids))
            if name.endswith(".rels"):
                base = "" if name == "_rels/.rels" else str(PurePosixPath(name).parent.parent)
                relationships = ET.fromstring(archive.read(name))
                identifiers = [node.get("Id") for node in relationships]
                self.assertEqual(len(identifiers), len(set(identifiers)), name)
                for relation in relationships:
                    self.assertNotEqual(relation.get("TargetMode"), "External")
                    target = posixpath.normpath(posixpath.join(base, relation.get("Target"))).lstrip("/")
                    self.assertIn(target, names, (name, target))

    def test_standard_library_export_preserves_manifest_order_image_bytes_and_page_size(self):
        result = self.build()
        self.assertEqual(result.returncode, 0, result.stderr)
        with ZipFile(self.output) as archive:
            self.assert_archive(archive)
            presentation = ET.fromstring(archive.read("ppt/presentation.xml"))
            size = presentation.find("p:sldSz", NS)
            self.assertEqual((int(size.get("cx")), int(size.get("cy"))), (12192000, 6858000))
            ids = presentation.findall("p:sldIdLst/p:sldId", NS)
            self.assertEqual(len(ids), len(self.order))
            rels = ET.fromstring(archive.read("ppt/_rels/presentation.xml.rels"))
            by_id = {node.get("Id"): node.get("Target") for node in rels}
            for identifier, expected in zip(ids, self.order):
                part = posixpath.normpath("ppt/" + by_id[identifier.get("{" + NS["r"] + "}id")])
                slide = ET.fromstring(archive.read(part))
                pictures = slide.findall("p:cSld/p:spTree/p:pic", NS)
                self.assertEqual(len(pictures), 1)
                picture = pictures[0]
                offset = picture.find("p:spPr/a:xfrm/a:off", NS)
                extent = picture.find("p:spPr/a:xfrm/a:ext", NS)
                self.assertEqual(offset.attrib, {"x": "0", "y": "0"})
                self.assertEqual(extent.attrib, {"cx": "12192000", "cy": "6858000"})
                relation_id = picture.find("p:blipFill/a:blip", NS).get("{" + NS["r"] + "}embed")
                relationships = ET.fromstring(archive.read(str(PurePosixPath(part).parent / "_rels" / (PurePosixPath(part).name + ".rels"))))
                target = next(node.get("Target") for node in relationships if node.get("Id") == relation_id)
                self.assertEqual(archive.read(posixpath.normpath(posixpath.join(posixpath.dirname(part), target))), self.images[expected])
            self.assertFalse(any(name.startswith("ppt/embeddings/") for name in archive.namelist()))
            self.assertEqual(sum(name.startswith("ppt/media/") for name in archive.namelist()), len(self.images))

    @unittest.skipUnless(importlib.util.find_spec("pptx"), "独立回读使用已有 python-pptx")
    def test_independent_reader_opens_slides_masters_layouts_and_original_images(self):
        result = self.build()
        self.assertEqual(result.returncode, 0, result.stderr)
        from pptx import Presentation
        presentation = Presentation(self.output)
        self.assertEqual(len(presentation.slides), 3)
        self.assertEqual(len(presentation.slide_masters), 1)
        self.assertEqual(len(presentation.slide_layouts), 1)
        for slide, expected in zip(presentation.slides, self.order):
            self.assertEqual(len(slide.shapes), 1)
            self.assertEqual(slide.shapes[0].image.blob, self.images[expected])
            self.assertEqual((slide.shapes[0].image.size), (16, 9))
            self.assertEqual(slide.slide_layout.name, "Blank")

    @unittest.skipUnless(importlib.util.find_spec("pptx") and importlib.util.find_spec("PIL") and importlib.util.find_spec("olefile"), "OLE 独立回读使用已有 python-pptx、Pillow 和 olefile")
    def test_html_ole_and_icon_remain_on_first_slide_with_original_html_bytes(self):
        html = self.directory / "原件 (讲义).html"
        original = ("<!doctype html><html><body>中文 & 原件</body></html>" * 150).encode("utf-8")
        html.write_bytes(original)
        result = self.build(html)
        self.assertEqual(result.returncode, 0, result.stderr)
        from pptx import Presentation
        import olefile
        presentation = Presentation(self.output)
        self.assertEqual([len(slide.shapes) for slide in presentation.slides], [2, 1, 1])
        shape = presentation.slides[0].shapes[1]
        self.assertEqual(shape.ole_format.prog_id, "Package")
        self.assertTrue(shape.ole_format.show_as_icon)
        with olefile.OleFileIO(io.BytesIO(shape.ole_format.blob)) as package:
            raw = package.openstream("\x01Ole10Native").read()
        position = 6
        label_end = raw.index(b"\0", position)
        self.assertEqual(raw[position:label_end].decode("gbk"), "原件讲义.html")
        position = raw.index(b"\0", label_end + 1) + 1 + 4
        position = raw.index(b"\0", position) + 1
        size = struct.unpack_from("<I", raw, position)[0]
        self.assertEqual(raw[position + 4:position + 4 + size], original)
        self.assertEqual(html.read_bytes(), original)
        with ZipFile(self.output) as archive:
            self.assert_archive(archive)
            slide = ET.fromstring(archive.read("ppt/slides/slide1.xml"))
            ole = slide.find(".//p:oleObj", NS)
            self.assertIsNotNone(ole.find("p:embed", NS))
            self.assertIsNotNone(ole.find("p:pic/p:blipFill/a:blip", NS))
            relationships = ET.fromstring(archive.read("ppt/slides/_rels/slide1.xml.rels"))
            relation = next(node for node in relationships if node.get("Id") == ole.get("{" + NS["r"] + "}id"))
            self.assertTrue(relation.get("Type").endswith("/oleObject"))
            self.assertEqual(sum(name.startswith("ppt/embeddings/") for name in archive.namelist()), 1)

    def test_empty_manifest_reports_error_without_creating_output(self):
        self.manifest([])
        result = self.build()
        self.assertEqual(result.returncode, 2)
        self.assertIn("截图", result.stderr)
        self.assertFalse(self.output.exists())

    @unittest.skipUnless(importlib.util.find_spec("pptx") and importlib.util.find_spec("PIL") and importlib.util.find_spec("olefile"), "OLE 独立回读使用已有 python-pptx、Pillow 和 olefile")
    def test_short_html_ole_attachment_survives_the_complete_export(self):
        html = self.directory / "tiny.html"
        original = b"<html>short</html>"
        html.write_bytes(original)
        result = self.build(html)
        self.assertEqual(result.returncode, 0, result.stderr)
        from pptx import Presentation
        import olefile
        presentation = Presentation(self.output)
        blob = presentation.slides[0].shapes[1].ole_format.blob
        with olefile.OleFileIO(io.BytesIO(blob), raise_defects=olefile.DEFECT_INCORRECT) as package:
            raw = package.openstream("\x01Ole10Native").read()
        self.assertTrue(raw.endswith(struct.pack("<I", len(original)) + original))
        self.assertEqual(struct.unpack_from("<I", raw)[0], len(raw) - 4)

    def test_invalid_screenshot_preserves_existing_output_and_cleans_temporary_file(self):
        previous = "已有导出结果".encode("utf-8")
        self.output.write_bytes(previous)
        (self.directory / "后页.png").write_bytes(b"not an image")
        result = self.build()
        self.assertEqual(result.returncode, 2)
        self.assertIn("PNG 或 JPEG", result.stderr)
        self.assertEqual(self.output.read_bytes(), previous)
        self.assertEqual(list(self.directory.glob(".pptx-*.tmp")), [])


if __name__ == "__main__":
    unittest.main()
