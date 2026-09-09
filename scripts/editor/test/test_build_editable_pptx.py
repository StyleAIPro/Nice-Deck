"""可编辑 PPTX 的原生对象、独立回读和标准库运行回归。"""
import copy
import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import test_build_pptx as baseline


SCRIPT = Path(__file__).resolve().parents[2] / "html2pptx" / "build_editable_pptx.py"
NS = baseline.NS


def paragraph(text, align="left", size=24):
    return {"align": align, "runs": [{"text": text, "fontFamily": "Arial", "fontSize": size, "color": "223344"}]}


def manifest():
    first = paragraph("可编辑 <标题> & 中文", "center", 48)
    first["lineHeight"] = 60
    first["runs"][0].update(fontFamily="Noto Sans CJK SC", bold=True, italic=True, underline=True)
    return {"schema": 1, "mode": "editable", "width": 1920, "height": 1080, "slides": [
        {"name": "第一 & 页", "warnings": ["复杂背景已栅格化"], "elements": [
            {"type": "text", "x": 120, "y": 80, "width": 1000, "height": 180,
             "paragraphs": [first, paragraph("第二段", "right", 28)]},
            {"type": "shape", "x": 100, "y": 300, "width": 300, "height": 80,
             "geometry": "roundRect", "radius": 20, "fill": "C51624", "opacity": 0.5,
             "border": {"color": "000000", "width": 2, "opacity": 0.25}},
            {"type": "shape", "x": 100, "y": 400, "width": 300, "height": 0,
             "geometry": "line", "border": {"color": "112233", "width": 3}},
            {"type": "shape", "x": 450, "y": 300, "width": 150, "height": 80, "geometry": "ellipse", "fill": "00AAFF"},
            {"type": "shape", "x": 620, "y": 300, "width": 150, "height": 80, "geometry": "rect"},
            {"type": "image", "x": 900, "y": 300, "width": 640, "height": 360, "file": "原图.png"},
            {"type": "image", "x": 1550, "y": 300, "width": 160, "height": 90, "file": "原图.jpg"},
            {"type": "table", "x": 120, "y": 760, "width": 800, "height": 180, "columns": [300, 500], "rows": [
                {"height": 60, "cells": [
                    {"paragraphs": [paragraph("名称")], "fill": "EEEEEE", "margin": 8, "border": {"color": "777777", "width": 2}},
                    {"paragraphs": [paragraph("数值", "center")], "fill": "EEEEEE"}]},
                {"height": 120, "cells": [{"paragraphs": [paragraph("收入")]}, {"paragraphs": [paragraph("42", "right")]}]},
            ]},
        ]},
        {"name": "第二页", "elements": [{"type": "image", "x": 0, "y": 0, "width": 1920, "height": 1080, "file": "原图.png"}]},
    ]}


class BuildEditablePptxTest(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.directory = Path(self.temporary.name)
        self.input = self.directory / "manifest"
        self.input.mkdir()
        self.output = self.directory / "输出 & 可编辑.pptx"
        (self.input / "原图.png").write_bytes(baseline.png())
        (self.input / "原图.jpg").write_bytes(baseline.JPEG)
        self.data = manifest()

    def build(self, embed=None):
        (self.input / "manifest.json").write_text(json.dumps(self.data, ensure_ascii=False), encoding="utf-8")
        args = [sys.executable] + (["-S"] if embed is None else []) + [str(SCRIPT), str(self.input), str(self.output)]
        if embed is not None:
            args.append(str(embed))
        return subprocess.run(args, capture_output=True, text=True)

    def test_standard_library_writes_native_objects_in_order_with_deduplicated_original_images(self):
        result = self.build()
        self.assertEqual(result.returncode, 0, result.stderr)
        with ZipFile(self.output) as archive:
            baseline.BuildPptxTest.assert_archive(self, archive)
            first = ET.fromstring(archive.read("ppt/slides/slide1.xml"))
            tree = first.find("p:cSld/p:spTree", NS)
            self.assertEqual([node.tag.rsplit("}", 1)[-1] for node in list(tree)[2:]], ["sp", "sp", "cxnSp", "sp", "sp", "pic", "pic", "graphicFrame"])
            self.assertEqual(first.find("p:cSld", NS).get("name"), "第一 & 页")
            ids = [node.get("id") for node in first.findall(".//p:cNvPr", NS)]
            self.assertEqual(len(ids), len(set(ids)))
            self.assertEqual(first.find(".//a:rPr", NS).get("sz"), "2400")
            self.assertEqual(first.find(".//a:lnSpc/a:spcPts", NS).get("val"), "3000")
            self.assertEqual(first.find(".//a:srgbClr[@val='C51624']/a:alpha", NS).get("val"), "50000")
            self.assertEqual(first.find(".//a:srgbClr[@val='000000']/a:alpha", NS).get("val"), "25000")
            self.assertEqual(first.find(".//a:gd[@name='adj']", NS).get("fmla"), "val 25000")
            self.assertEqual(sum(name.startswith("ppt/media/") for name in archive.namelist()), 2)
            self.assertIn(baseline.png(), [archive.read(name) for name in archive.namelist() if name.startswith("ppt/media/")])

    @unittest.skipUnless(importlib.util.find_spec("pptx"), "独立回读需要已有 python-pptx；不是 writer 运行依赖")
    def test_independent_reader_edits_rich_text_shapes_and_native_table_then_saves(self):
        result = self.build()
        self.assertEqual(result.returncode, 0, result.stderr)
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE, MSO_AUTO_SHAPE_TYPE
        from pptx.enum.text import PP_ALIGN
        deck = Presentation(self.output)
        self.assertEqual((deck.slide_width, deck.slide_height), (12192000, 6858000))
        self.assertEqual([len(slide.shapes) for slide in deck.slides], [8, 1])
        shapes = deck.slides[0].shapes
        text = shapes[0]
        self.assertEqual(text.shape_type, MSO_SHAPE_TYPE.TEXT_BOX)
        self.assertEqual((text.left, text.top, text.width, text.height), (762000, 508000, 6350000, 1143000))
        self.assertEqual(text.text, "可编辑 <标题> & 中文\n第二段")
        first = text.text_frame.paragraphs[0]
        self.assertEqual(first.alignment, PP_ALIGN.CENTER)
        self.assertEqual(first.line_spacing.pt, 30)
        font = first.runs[0].font
        self.assertEqual((font.name, font.size.pt, str(font.color.rgb)), ("Noto Sans CJK SC", 24, "223344"))
        self.assertTrue(font.bold and font.italic and font.underline)
        self.assertEqual(shapes[1].auto_shape_type, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE)
        self.assertEqual(shapes[2].shape_type, MSO_SHAPE_TYPE.LINE)
        self.assertEqual(shapes[3].auto_shape_type, MSO_AUTO_SHAPE_TYPE.OVAL)
        self.assertEqual(shapes[4].auto_shape_type, MSO_AUTO_SHAPE_TYPE.RECTANGLE)
        self.assertEqual(shapes[5].image.blob, baseline.png())
        self.assertEqual(shapes[6].image.blob, baseline.JPEG)
        table = shapes[7].table
        self.assertEqual([[cell.text for cell in row.cells] for row in table.rows], [["名称", "数值"], ["收入", "42"]])
        self.assertEqual([column.width for column in table.columns], [1905000, 3175000])
        self.assertEqual([row.height for row in table.rows], [381000, 762000])
        self.assertEqual(table.cell(0, 0).margin_left, 50800)
        self.assertEqual(str(table.cell(0, 0).fill.fore_color.rgb), "EEEEEE")
        first.runs[0].text = "已编辑标题"
        table.cell(1, 1).text = "84"
        deck.save(self.output)
        changed = Presentation(self.output)
        self.assertTrue(changed.slides[0].shapes[0].text.startswith("已编辑标题"))
        self.assertEqual(changed.slides[0].shapes[7].table.cell(1, 1).text, "84")

    @unittest.skipUnless(importlib.util.find_spec("pptx") and importlib.util.find_spec("PIL"), "附件独立回读需要已有 python-pptx 和 Pillow")
    def test_original_html_attachment_uses_unique_ids_and_only_the_first_slide(self):
        html = self.directory / "原始 HTML.html"
        original = b"<!doctype html><h1>editable source</h1>"
        html.write_bytes(original)
        result = self.build(html)
        self.assertEqual(result.returncode, 0, result.stderr)
        from pptx import Presentation
        deck = Presentation(self.output)
        self.assertEqual([len(slide.shapes) for slide in deck.slides], [9, 1])
        attachment = deck.slides[0].shapes[-1]
        self.assertEqual(attachment.ole_format.prog_id, "Package")
        self.assertIn(original, attachment.ole_format.blob)
        self.assertEqual(html.read_bytes(), original)
        self.assertFalse((self.input / "_ole_icon.png").exists())
        with ZipFile(self.output) as archive:
            baseline.BuildPptxTest.assert_archive(self, archive)
            first = ET.fromstring(archive.read("ppt/slides/slide1.xml"))
            ids = [node.get("id") for node in first.findall(".//p:cNvPr", NS)]
            self.assertEqual(len(ids), len(set(ids)))

    def test_tracking_opacity_fixed_lines_and_cell_insets_survive_native_readback(self):
        text = self.data["slides"][0]["elements"][0]
        text.update(wrap=False, topInset=6)
        text["paragraphs"][0]["runs"][0].update(charSpacing=-2, opacity=0.4)
        text["paragraphs"][0]["runs"].append({
            "text": " 第二片段\n软换行", "fontFamily": "Arial", "fontSize": 24,
            "color": "336699", "charSpacing": 2,
        })
        cell = self.data["slides"][0]["elements"][-1]["rows"][0]["cells"][0]
        cell.update(wrap=False, marginTop=2, marginRight=4, marginBottom=6, marginLeft=8, verticalAlign="center",
                    borderTop={"color": "FF0000", "width": 4, "opacity": 0.5}, borderRight=None,
                    borderBottom={"color": "00FF00", "width": 6})
        result = self.build()
        self.assertEqual(result.returncode, 0, result.stderr)
        with ZipFile(self.output) as archive:
            first = ET.fromstring(archive.read("ppt/slides/slide1.xml"))
            body = first.find(".//p:txBody/a:bodyPr", NS)
            self.assertEqual((body.get("wrap"), body.get("tIns")), ("none", "38100"))
            runs = first.findall(".//p:txBody/a:p/a:r/a:rPr", NS)
            self.assertEqual([run.get("spc") for run in runs[:3]], ["-100", "100", "100"])
            self.assertEqual(runs[0].find("a:solidFill/a:srgbClr/a:alpha", NS).get("val"), "40000")
            self.assertEqual(first.find(".//p:txBody/a:p/a:pPr/a:lnSpc/a:spcPts", NS).get("val"), "3000")
            self.assertEqual(len(first.findall(".//p:txBody/a:p/a:br", NS)), 1)
            self.assertEqual([body.get("wrap") for body in first.findall(".//a:tc/a:txBody/a:bodyPr", NS)], ["none", "square", "square", "square"])
            properties = first.find(".//a:tc/a:tcPr", NS)
            self.assertEqual({key: properties.get(key) for key in ("marL", "marR", "marT", "marB", "anchor")}, {
                "marL": "50800", "marR": "25400", "marT": "12700", "marB": "38100", "anchor": "ctr",
            })
            self.assertEqual(properties.find("a:lnT/a:solidFill/a:srgbClr", NS).get("val"), "FF0000")
            self.assertEqual(properties.find("a:lnT/a:solidFill/a:srgbClr/a:alpha", NS).get("val"), "50000")
            self.assertEqual(properties.find("a:lnT", NS).get("w"), "25400")
            self.assertIsNotNone(properties.find("a:lnR/a:noFill", NS))
            self.assertEqual(properties.find("a:lnB/a:solidFill/a:srgbClr", NS).get("val"), "00FF00")
            self.assertEqual(properties.find("a:lnB", NS).get("w"), "38100")
            self.assertEqual(properties.find("a:lnL/a:solidFill/a:srgbClr", NS).get("val"), "777777")
            self.assertEqual(properties.find("a:lnL", NS).get("w"), "12700")
        if importlib.util.find_spec("pptx"):
            from pptx import Presentation
            from pptx.enum.text import MSO_VERTICAL_ANCHOR
            slide = Presentation(self.output).slides[0]
            self.assertFalse(slide.shapes[0].text_frame.word_wrap)
            self.assertEqual(slide.shapes[0].text_frame.margin_top, 38100)
            self.assertEqual(slide.shapes[0].text_frame.paragraphs[0].text, "可编辑 <标题> & 中文 第二片段\v软换行")
            table_cell = slide.shapes[-1].table.cell(0, 0)
            self.assertFalse(table_cell.text_frame.word_wrap)
            self.assertTrue(slide.shapes[-1].table.cell(0, 1).text_frame.word_wrap)
            self.assertEqual(table_cell.vertical_anchor, MSO_VERTICAL_ANCHOR.MIDDLE)
            self.assertEqual((table_cell.margin_top, table_cell.margin_right, table_cell.margin_bottom, table_cell.margin_left), (12700, 25400, 38100, 50800))

    def test_zip_write_failure_preserves_previous_output_and_removes_staging_file(self):
        result = self.build()
        self.assertEqual(result.returncode, 0, result.stderr)
        previous = self.output.read_bytes()
        spec = importlib.util.spec_from_file_location("editable_writer", SCRIPT)
        writer = importlib.util.module_from_spec(spec)
        sys.path.insert(0, str(SCRIPT.parent))
        try:
            spec.loader.exec_module(writer)
        finally:
            sys.path.pop(0)
        with patch.object(writer.ZipFile, "writestr", side_effect=OSError("模拟写入失败")):
            with self.assertRaisesRegex(OSError, "模拟写入失败"):
                writer.build_editable_pptx(self.input, self.output)
        self.assertEqual(self.output.read_bytes(), previous)
        self.assertEqual(list(self.directory.glob(".pptx-*.tmp")), [])

    def test_invalid_objects_fail_without_replacing_an_existing_output(self):
        invalid = [
            {"type": "unknown", "x": 0, "y": 0, "width": 1, "height": 1},
            {"type": "shape", "geometry": "rect", "x": float("nan"), "y": 0, "width": 1, "height": 1},
            {"type": "shape", "geometry": "rect", "x": 0, "y": 0, "width": 1, "height": 1, "fill": "not-hex"},
            {"type": "shape", "geometry": "rect", "x": 0, "y": 0, "width": 1, "height": 1, "opacity": 2},
            {"type": "text", "x": 0, "y": 0, "width": 100, "height": 100, "paragraphs": [paragraph("坏\0文本")]},
        ]
        malformed_table = copy.deepcopy(self.data["slides"][0]["elements"][-1])
        malformed_table["rows"][0]["cells"].pop()
        invalid.append(malformed_table)
        malformed_text = copy.deepcopy(self.data["slides"][0]["elements"][0])
        malformed_text["paragraphs"][0]["runs"][0].pop("color")
        invalid.append(malformed_text)
        malformed_text = copy.deepcopy(self.data["slides"][0]["elements"][0])
        malformed_text["paragraphs"][0]["align"] = []
        invalid.append(malformed_text)
        previous = b"keep existing export"
        for element in invalid:
            with self.subTest(element=element):
                self.output.write_bytes(previous)
                self.data = manifest()
                self.data["slides"][1]["elements"] = [element]
                result = self.build()
                self.assertEqual(result.returncode, 2, result.stderr)
                self.assertIn("可编辑 PPTX 组装失败", result.stderr)
                self.assertEqual(self.output.read_bytes(), previous)
                self.assertEqual(list(self.directory.glob(".pptx-*.tmp")), [])

    def test_media_paths_cannot_escape_the_manifest_directory(self):
        external = self.directory / "outside.png"
        external.write_bytes(baseline.png())
        for file in ["../outside.png", str(external), "nested/../../outside.png", "C:\\outside.png"]:
            with self.subTest(file=file):
                self.data["slides"][1]["elements"][0]["file"] = file
                result = self.build()
                self.assertEqual(result.returncode, 2, result.stderr)
                self.assertIn("可编辑 PPTX 组装失败", result.stderr)
                self.assertFalse(self.output.exists())
        if sys.platform != "win32":
            (self.input / "linked.png").symlink_to(external)
            self.data["slides"][1]["elements"][0]["file"] = "linked.png"
            self.assertEqual(self.build().returncode, 2)
        self.assertEqual(external.read_bytes(), baseline.png())


if __name__ == "__main__":
    unittest.main()
